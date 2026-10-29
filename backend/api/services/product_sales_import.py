"""Strict, header-tolerant parser for exported Union product sales files.

Only sales fields are kept, never original file bytes. A filename must contain
an unambiguous product ID that is linked to an event ticket type in the site.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from pathlib import Path

from fastapi import HTTPException

MAX_UPLOAD_BYTES = 6 * 1024 * 1024
MAX_ROWS = 5000


def product_ids_in_filename(filename: str) -> list[int]:
    stem = Path(filename or '').stem
    named = re.findall(r'(?i)product\s*(?:id|no|number|#)?\s*[-_ #(:]*\s*(\d{2,11})(?![a-z0-9])', stem)
    if named:
        return list(dict.fromkeys(int(p) for p in named))
    return list(dict.fromkeys(int(p) for p in re.findall(r'(?<!\d)\d{4,11}(?!\d)', stem)))


def canonical_header(value: object) -> str:
    text = unicodedata.normalize('NFKD', str(value or ''))
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', text.casefold())


ALIASES = {
    'cid': ('cid', 'customercid', 'studentcid', 'membercid', 'studentnumber', 'collegeid', 'imperialid'),
    'first_name': ('firstname', 'first', 'forename', 'givenname', 'customerfirstname', 'memberfirstname'),
    'last_name': ('lastname', 'last', 'surname', 'familyname', 'customersurname', 'customerlastname'),
    'full_name': ('name', 'fullname', 'customername', 'purchasername', 'buyername', 'membername'),
    'email': ('email', 'emailaddress', 'customeremail', 'buyeremail'),
    'order_number': ('orderno', 'ordernumber', 'orderid', 'transactionid', 'transactionnumber', 'receiptno', 'saleid'),
    'quantity': ('quantity', 'qty', 'quantitysold', 'items', 'numberofitems', 'tickets'),
    'product_line_id': ('productlineid', 'lineid', 'itemlineid', 'productlineno'),
}


def _data_rows(filename: str, content: bytes) -> list[list[object]]:
    suffix = Path(filename).suffix.lower()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, 'The sales export must be 6 MB or smaller')
    if suffix == '.xlsx':
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows = []
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i > MAX_ROWS + 30:
                        raise HTTPException(413, 'Too many rows (maximum 5000)')
                    rows.append(list(row))
                return rows
            finally:
                workbook.close()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(400, 'Unable to read XLSX sales export') from exc
    if suffix == '.xls':
        try:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content)
            sheet = workbook.sheet_by_index(0)
            if sheet.nrows > MAX_ROWS + 30:
                raise HTTPException(413, 'Too many rows (maximum 5000)')
            return [sheet.row_values(i) for i in range(sheet.nrows)]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(400, 'Unable to read XLS sales export') from exc
    if suffix != '.csv':
        raise HTTPException(400, 'Upload a CSV, XLS, or XLSX product sales export')
    try:
        encoding = 'utf-16' if content.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            text = content.decode('cp1252')
        sample = text[:8192]
        # The export can begin with title/date lines that have no delimiter.
        # Sniffer frequently guesses wrong on those, so score the first lines.
        first_lines = sample.splitlines()[:25]
        separators = (',', ';', '\t', '|')
        delimiter = max(separators, key=lambda char: sum(line.count(char) for line in first_lines))
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = []
        for i, row in enumerate(reader):
            if i > MAX_ROWS + 30:
                raise HTTPException(413, 'Too many rows (maximum 5000)')
            rows.append(row)
        return rows
    except HTTPException:
        raise
    except (ValueError, csv.Error) as exc:
        raise HTTPException(400, 'Unable to read CSV sales export') from exc


def parse_export(filename: str, content: bytes) -> list[dict]:
    lines = _data_rows(filename, content)
    header_index = None
    columns = {}
    for i, fields in enumerate(lines[:25]):
        keys = [canonical_header(cell) for cell in fields]
        found = {field: next((j for j, key in enumerate(keys) if key in aliases), None)
                 for field, aliases in ALIASES.items()}
        if (found['cid'] is not None or found['full_name'] is not None or
                found['first_name'] is not None and found['last_name'] is not None):
            header_index = i
            columns = found
            break
    if header_index is None:
        raise HTTPException(400, 'No recognised CID or buyer name columns. Expected CID or first/last name (or customer name).')
    result = []
    for i, cells in enumerate(lines[header_index + 1:], start=header_index + 2):
        def field(name: str) -> str:
            index = columns[name]
            if index is None or index >= len(cells) or cells[index] is None:
                return ''
            value = cells[index]
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()
        cid, first, last, full = field('cid'), field('first_name'), field('last_name'), field('full_name')
        name = full or ' '.join(v for v in (first, last) if v)
        if full and not first and not last and ',' in full:
            surname, given = full.split(',', 1)
            if surname.strip() and given.strip():
                name = f'{given.strip()} {surname.strip()}'
        if not any((cid, name, field('order_number'))):
            continue
        quantity = field('quantity')
        try:
            number = int(float(quantity)) if quantity else 1
            if not 1 <= number <= 100 or number != float(quantity or 1):
                raise ValueError()
        except ValueError as exc:
            raise HTTPException(400, f'Invalid quantity on row {i}') from exc
        line = field('product_line_id')
        try:
            line_id = int(float(line)) if line else None
        except ValueError as exc:
            raise HTTPException(400, f'Invalid product line ID on row {i}') from exc
        # Buyer with no CID or name cannot be linked safely to a reservation.
        if not cid and not name:
            raise HTTPException(400, f'Row {i} has no CID or buyer name')
        if len(cid) > 40 or len(name) > 255 or len(field('order_number')) > 100 or len(field('email')) > 255:
            raise HTTPException(400, f'One or more fields on row {i} are too long')
        result.append({'source_row': i, 'cid': cid or None, 'buyer_name': name or None,
                       'email': field('email') or None, 'order_number': field('order_number') or None,
                       'quantity': number, 'product_line_id': line_id})
    if not result:
        raise HTTPException(400, 'No sales records found in the export')
    if len(result) > MAX_ROWS:
        raise HTTPException(413, 'Too many sales records')
    return result
